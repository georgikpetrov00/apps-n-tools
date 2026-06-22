#!/usr/bin/env python3
"""
SAP portfolio average-price calculator.

Reads PortfolioDetails.xlsx (the SAP Equity Plan portfolio export) and prints
four progressively-inclusive weighted average prices per share:

  1. Avg price of BOUGHT shares only (rows with Contribution type == "Purchase")
  2. Avg price of BOUGHT + COMPANY MATCH shares
  3. Avg price of BOUGHT + COMPANY MATCH + ALREADY-RECEIVED awards
       (Award rows where Available quantity > 0 -- i.e. already vested)
  4. Avg price of BOUGHT + COMPANY MATCH + ALL awards (vested + still-locked)

Free shares (Company match, Awards) contribute QUANTITY but ZERO cost to the
weighted average -- that is, "what did I actually pay per share I now hold".

Runs identically on Windows and macOS:
    python portfolio_averages.py
    python portfolio_averages.py path/to/PortfolioDetails.xlsx

Requires: openpyxl  (install once with:  pip install openpyxl)
"""

from __future__ import annotations

import sys
import warnings
from pathlib import Path
from datetime import datetime

# Silence openpyxl's harmless "workbook contains no default style" warning that
# the SAP export triggers -- the file is fine, it just lacks the optional block.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

try:
    import openpyxl
except ImportError:
    print("ERROR: this script needs the 'openpyxl' package.")
    print("Install it with:  pip install openpyxl")
    sys.exit(1)


SHEET_NAME = "Portfolio details"
HEADER_ROW = 7  # row containing column headers in the SAP export


def find_file(arg: str | None) -> Path:
    """Locate the portfolio XLSX. Accept an explicit path, else look next to this script."""
    if arg:
        p = Path(arg)
        if not p.exists():
            print(f"ERROR: file not found: {p}")
            sys.exit(1)
        return p
    here = Path(__file__).resolve().parent
    default = here / "PortfolioDetails.xlsx"
    if default.exists():
        return default
    # Also try current working directory as a fallback.
    cwd_candidate = Path.cwd() / "PortfolioDetails.xlsx"
    if cwd_candidate.exists():
        return cwd_candidate
    print("ERROR: PortfolioDetails.xlsx not found next to the script or in the current folder.")
    print("Pass the path explicitly:  python portfolio_averages.py /path/to/PortfolioDetails.xlsx")
    sys.exit(1)


def load_rows(xlsx_path: Path) -> list[dict]:
    """Read the data rows from the 'Portfolio details' sheet into a list of dicts."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        print(f"ERROR: sheet '{SHEET_NAME}' not found. Sheets present: {wb.sheetnames}")
        sys.exit(1)
    ws = wb[SHEET_NAME]

    headers = [c.value for c in ws[HEADER_ROW]]
    rows: list[dict] = []
    for raw in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        # Skip fully empty trailing rows.
        if all(v is None for v in raw):
            continue
        row = dict(zip(headers, raw))
        rows.append(row)
    return rows


def num(v) -> float:
    """Coerce a cell value to float; None / blanks become 0."""
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def weighted_avg(total_cost: float, total_qty: float) -> float:
    return total_cost / total_qty if total_qty > 0 else 0.0


def classify(row: dict) -> str:
    """
    Bucket a row into one of:
      - 'purchase'           -> shares you paid for
      - 'match'              -> Company match shares (free)
      - 'award_vested'       -> Award rows already available (Available quantity > 0)
      - 'award_unvested'     -> Award rows still locked (Available quantity == 0)
      - 'other'              -> anything we don't recognize; ignored in the averages
    """
    ctype = (row.get("Contribution type") or "").strip().lower()
    available_qty = num(row.get("Available quantity"))

    if ctype == "purchase":
        return "purchase"
    if ctype == "company match":
        return "match"
    if ctype == "award":
        return "award_vested" if available_qty > 0 else "award_unvested"
    return "other"


def main() -> None:
    arg = sys.argv[1] if len(sys.argv) > 1 else None
    xlsx_path = find_file(arg)

    rows = load_rows(xlsx_path)

    # Each bucket tracks (total_qty, total_cost_paid).
    # For purchases:   cost = strike_price * outstanding_qty   (what you actually paid)
    # For free shares: cost = 0, only quantity is added.
    # We use "Outstanding quantity" because it reflects what's still in your account
    # (sold/released shares are no longer relevant to your current avg cost).
    buckets = {
        "purchase":        {"qty": 0.0, "cost": 0.0, "count": 0},
        "match":           {"qty": 0.0, "cost": 0.0, "count": 0},
        "award_vested":    {"qty": 0.0, "cost": 0.0, "count": 0},
        "award_unvested":  {"qty": 0.0, "cost": 0.0, "count": 0},
    }

    skipped_other = 0
    for row in rows:
        bucket = classify(row)
        if bucket == "other":
            skipped_other += 1
            continue

        qty = num(row.get("Outstanding quantity"))
        if qty <= 0:
            # Nothing left in this lot -- can't meaningfully add to the avg.
            continue

        if bucket == "purchase":
            cost_basis = num(row.get("Strike price / Cost basis"))
            cost = cost_basis * qty
        else:
            # Free shares: quantity counts, cost is zero.
            cost = 0.0

        buckets[bucket]["qty"] += qty
        buckets[bucket]["cost"] += cost
        buckets[bucket]["count"] += 1

    # Progressive sums for the four averages.
    p = buckets["purchase"]
    m = buckets["match"]
    av = buckets["award_vested"]
    au = buckets["award_unvested"]

    avg1_qty,  avg1_cost  = p["qty"], p["cost"]
    avg2_qty,  avg2_cost  = avg1_qty + m["qty"],  avg1_cost + m["cost"]
    avg3_qty,  avg3_cost  = avg2_qty + av["qty"], avg2_cost + av["cost"]
    avg4_qty,  avg4_cost  = avg3_qty + au["qty"], avg3_cost + au["cost"]

    avg1 = weighted_avg(avg1_cost, avg1_qty)
    avg2 = weighted_avg(avg2_cost, avg2_qty)
    avg3 = weighted_avg(avg3_cost, avg3_qty)
    avg4 = weighted_avg(avg4_cost, avg4_qty)

    # ------------------------------------------------------------------
    # Build the report.
    # ------------------------------------------------------------------
    lines: list[str] = []
    w = lines.append

    def render_table(headers: list[str], rows: list[list[str]],
                     aligns: list[str] | None = None) -> list[str]:
        """
        Render a simple bordered ASCII table.
        aligns: per-column 'l' or 'r' (defaults to 'l' for col 0, 'r' for the rest).
        """
        if aligns is None:
            aligns = ["l"] + ["r"] * (len(headers) - 1)
        col_widths = [
            max(len(headers[i]), *(len(r[i]) for r in rows)) if rows else len(headers[i])
            for i in range(len(headers))
        ]

        def fmt_cell(text: str, width: int, align: str) -> str:
            return text.ljust(width) if align == "l" else text.rjust(width)

        def border(left: str, mid: str, right: str, fill: str = "-") -> str:
            return left + mid.join(fill * (w + 2) for w in col_widths) + right

        def row_line(cells: list[str]) -> str:
            inner = " | ".join(
                fmt_cell(c, col_widths[i], aligns[i]) for i, c in enumerate(cells)
            )
            return "| " + inner + " |"

        out = [
            border("+", "+", "+"),
            row_line(headers),
            border("+", "+", "+", "="),
        ]
        for r in rows:
            out.append(row_line(r))
        out.append(border("+", "+", "+"))
        return out

    def fmt_eur(x: float) -> str:
        return f"EUR {x:,.2f}"

    def fmt_qty(x: float) -> str:
        return f"{x:,.4f}"

    w("=" * 72)
    w("  SAP PORTFOLIO -- AVERAGE PRICE PER SHARE")
    w("=" * 72)
    w(f"  Source file : {xlsx_path}")
    w(f"  Generated   : {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    w(f"  Method      : weighted average (cost paid / shares held)")
    w(f"                Free shares (match, awards) add quantity at 0 cost.")
    w(f"                Uses 'Outstanding quantity' from each lot.")
    w("")

    # -- Bucket breakdown table --------------------------------------------
    w("Bucket breakdown")
    bucket_rows = [
        ["Purchase (you paid)",         str(p["count"]),  fmt_qty(p["qty"]),
         fmt_eur(p["cost"]),  fmt_eur(weighted_avg(p["cost"],  p["qty"]))],
        ["Company match (free)",        str(m["count"]),  fmt_qty(m["qty"]),
         fmt_eur(m["cost"]),  fmt_eur(weighted_avg(m["cost"],  m["qty"]))],
        ["Awards already received",     str(av["count"]), fmt_qty(av["qty"]),
         fmt_eur(av["cost"]), fmt_eur(weighted_avg(av["cost"], av["qty"]))],
        ["Awards still to be received", str(au["count"]), fmt_qty(au["qty"]),
         fmt_eur(au["cost"]), fmt_eur(weighted_avg(au["cost"], au["qty"]))],
    ]
    for line in render_table(
        ["Bucket", "Rows", "Quantity", "Total cost", "Avg / share"],
        bucket_rows,
    ):
        w(line)
    w("")

    # -- Four progressive averages table -----------------------------------
    w("Four progressive averages")
    avg_rows = [
        ["1", "Bought shares only",
         fmt_qty(avg1_qty), fmt_eur(avg1_cost), fmt_eur(avg1)],
        ["2", "Bought + company match",
         fmt_qty(avg2_qty), fmt_eur(avg2_cost), fmt_eur(avg2)],
        ["3", "Bought + match + already-received awards",
         fmt_qty(avg3_qty), fmt_eur(avg3_cost), fmt_eur(avg3)],
        ["4", "Bought + match + all awards (incl. unvested)",
         fmt_qty(avg4_qty), fmt_eur(avg4_cost), fmt_eur(avg4)],
    ]
    for line in render_table(
        ["#", "Scope", "Total shares", "Total cost", "Avg / share"],
        avg_rows,
        aligns=["r", "l", "r", "r", "r"],
    ):
        w(line)
    w("")

    if skipped_other:
        w(f"(Skipped {skipped_other} row(s) with unrecognized contribution type.)")
        w("")
    w("=" * 72)

    report = "\n".join(lines)
    print(report)

    # Save report next to the script, with today's date appended to the filename.
    date_suffix = datetime.now().strftime("%d-%m-%Y")
    out_path = Path(__file__).resolve().parent / f"portfolio_averages_report_{date_suffix}.txt"
    out_path.write_text(report + "\n", encoding="utf-8")
    print(f"\nReport saved to: {out_path}")


if __name__ == "__main__":
    main()
